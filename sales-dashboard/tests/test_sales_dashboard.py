"""Offline tests: python3 -m pytest -q  (needs pytest and openpyxl; no network)."""
import csv
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))
import make_sample_data  # noqa: E402
import sales_dashboard as sd  # noqa: E402

SAMPLE = HERE / "input" / "sales_2025.csv"
HEADER = ",".join(sd.REQUIRED)


def write_csv(tmp_path, lines, name="s.csv"):
    p = tmp_path / name
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def test_sample_file_matches_generator():
    """The committed CSV is exactly what make_sample_data.py produces (synthetic, reproducible)."""
    with open(SAMPLE, newline="") as f:
        committed = list(csv.DictReader(f))
    generated = make_sample_data.generate()
    assert len(committed) == len(generated)
    assert committed[0] == {k: str(v) for k, v in generated[0].items()}
    assert committed[-1] == {k: str(v) for k, v in generated[-1].items()}


def test_load_parses_types(tmp_path):
    p = write_csv(tmp_path, [HEADER, "SO-1,2025-03-04,East,Online,Tea,Green Tea Tin,2,$11.00,\"$1,022.00\""])
    rows, skipped = sd.load_sales(p)
    assert skipped == []
    r = rows[0]
    assert r["order_date"].month == 3 and r["units"] == 2
    assert r["unit_price"] == 11.0 and r["revenue"] == 1022.0


def test_headers_are_case_and_space_insensitive(tmp_path):
    header = ",".join(f" {c.upper()} " for c in sd.REQUIRED)
    p = write_csv(tmp_path, [header, "SO-1,2025-01-02,East,Online,Tea,Tin,1,1,1"])
    rows, _ = sd.load_sales(p)
    assert len(rows) == 1


def test_missing_column_is_an_error(tmp_path):
    p = write_csv(tmp_path, ["order_id,order_date,region", "SO-1,2025-01-01,East"])
    with pytest.raises(ValueError, match="missing column"):
        sd.load_sales(p)


def test_bad_rows_are_skipped_and_reported(tmp_path):
    p = write_csv(tmp_path, [
        HEADER,
        "SO-1,2025-01-05,East,Online,Tea,Tin,1,10,10",
        "SO-2,05/01/2025,East,Online,Tea,Tin,1,10,10",   # wrong date format
        "SO-3,2025-01-06,East,Online,Tea,Tin,two,10,10",  # not a number
        "SO-4,2025-01-07,,Online,Tea,Tin,1,10,10",        # blank region
        ",,,,,,,,",                                       # blank line: ignored, not an error
    ])
    rows, skipped = sd.load_sales(p)
    assert [r["order_id"] for r in rows] == ["SO-1"]
    assert [line for line, _ in skipped] == [3, 4, 5]
    assert "order_date" in skipped[0][1]


def test_kpis_match_an_independent_sum():
    rows, _ = sd.load_sales(SAMPLE)
    s = sd.summarize(rows)
    with open(SAMPLE, newline="") as f:
        raw = list(csv.DictReader(f))
    total = round(sum(float(r["revenue"]) for r in raw), 2)
    assert s["kpis"]["revenue"] == total
    assert s["kpis"]["orders"] == len(raw)
    assert s["kpis"]["units"] == sum(int(r["units"]) for r in raw)
    assert s["kpis"]["avg_order"] == round(total / len(raw), 2)


def test_breakdowns_add_up_to_the_total():
    s = sd.summarize(sd.load_sales(SAMPLE)[0])
    total = s["kpis"]["revenue"]
    assert len(s["monthly"]) == 12
    assert sum(s["monthly"]) == pytest.approx(total, abs=0.05)
    assert sum(v for _, v in s["regions"]) == pytest.approx(total, abs=0.05)
    assert sum(v for _, v in s["categories"]) == pytest.approx(total, abs=0.05)
    for i, m in enumerate(s["monthly"]):
        assert sum(ch[i] for ch in s["channels"].values()) == pytest.approx(m, abs=0.05)
    assert [v for _, v in s["regions"]] == sorted((v for _, v in s["regions"]), reverse=True)
    assert s["kpis"]["best_month_revenue"] == max(s["monthly"])


def test_months_with_no_sales_are_kept_as_zero(tmp_path):
    p = write_csv(tmp_path, [HEADER, "A,2025-01-10,E,Online,T,P,1,5,5", "B,2025-03-10,E,Online,T,P,1,7,7"])
    s = sd.summarize(sd.load_sales(p)[0])
    assert s["month_labels"] == ["Jan", "Feb", "Mar"]
    assert s["monthly"] == [5.0, 0.0, 7.0]


def test_nice_ticks_and_money():
    assert sd.nice_ticks(9409.5) == [0, 2500, 5000, 7500, 10000]
    assert sd.nice_ticks(6548.5) == [0, 2000, 4000, 6000, 8000]
    assert sd.nice_ticks(0) == [0, 1]
    assert sd.money(83778.5) == "$83,778.50"
    assert sd.money(28417.5, compact=True) == "$28.4K"
    assert sd.money(2000, compact=True) == "$2K"
    assert sd.money(950, compact=True) == "$950"


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    assert sd.main([str(SAMPLE), "--out", str(out)]) == 0
    return out


def test_html_is_self_contained_with_four_charts(built):
    page = (built / "dashboard.html").read_text()
    assert "<script" not in page and "<link" not in page and "@import" not in page
    assert not re.search(r'(src|href)="https?://', page)
    svgs = re.findall(r"<svg.*?</svg>", page, flags=re.S)
    assert len(svgs) == 4
    for svg in svgs:
        ET.fromstring(svg)  # each chart is well-formed XML
    assert page.count("Show data table") == 4  # every chart has a table twin
    assert "$83,778.50" in page and "1,656" in page


def test_html_supports_light_and_dark(built):
    page = (built / "dashboard.html").read_text()
    assert "prefers-color-scheme:dark" in page and ':root[data-theme="dark"]' in page


def test_chart_labels_stay_inside_the_chart(built):
    """Direct labels are placed with room to spare; none should run past the SVG's width."""
    page = (built / "dashboard.html").read_text()
    for svg in re.findall(r"<svg.*?</svg>", page, flags=re.S):
        width = float(re.search(r'viewBox="0 0 ([\d.]+)', svg).group(1))
        for x, text in re.findall(r'<text class="label" x="([\d.]+)"[^>]*>([^<]*)<', svg):
            assert float(x) + sd.text_width(text) <= width, text


def test_labels_are_html_escaped(tmp_path):
    p = write_csv(tmp_path, [HEADER, 'A,2025-01-10,<b>North</b>,Online,T&C,P,1,5,5'])
    page = sd.render_html(sd.summarize(sd.load_sales(p)[0]))
    assert "<b>North</b>" not in page and "&lt;b&gt;North&lt;/b&gt;" in page
    assert "T&amp;C" in page


def test_xlsx_has_four_native_charts(built):
    xlsx = built / "sales_report.xlsx"
    with zipfile.ZipFile(xlsx) as z:
        charts = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
        assert len(charts) == 4
        xml = [z.read(n).decode() for n in charts]
    kinds = [re.search(r"<(?:c:)?(lineChart|barChart)>", x).group(1) for x in xml]
    assert kinds.count("lineChart") == 2 and kinds.count("barChart") == 2
    for x in xml:
        assert '<delete val="1"' not in x and "<delete val=\"0\"" in x  # axes are visible
        assert "2A78D6" in x  # series use the validated palette


def test_xlsx_values_match_summary(built):
    from openpyxl import load_workbook
    wb = load_workbook(built / "sales_report.xlsx")
    assert wb.sheetnames == ["Summary", "Orders"]
    ws = wb["Summary"]
    s = sd.summarize(sd.load_sales(SAMPLE)[0])
    assert ws["B4"].value == s["kpis"]["revenue"]
    assert ws["A12"].value == "Jan" and ws["D23"].value == s["monthly"][-1]
    orders = wb["Orders"]
    assert orders.max_row == s["kpis"]["orders"] + 1
    assert "Orders" in orders.tables


def test_no_xlsx_flag_needs_only_the_standard_library(tmp_path):
    assert sd.main([str(SAMPLE), "--out", str(tmp_path), "--no-xlsx"]) == 0
    assert (tmp_path / "dashboard.html").exists()
    assert not (tmp_path / "sales_report.xlsx").exists()
